"""Read-only source audit and candidate extraction of newly supplied Rosstat prices.

Does not modify source workbooks, existing application datasets, or Git state.
Run with the bundled Python runtime; prices are levels, not official CPI indices.
"""

from __future__ import annotations

import argparse
import collections
import hashlib
import json
import math
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PRICE_FILES = [
    "sred_potreb_cen_2022.xlsx", "sred_potreb_cen_2023.xlsx",
    "sred_potreb_cen_2024.xlsx", "sred_potreb_cen_2025.xlsx",
    "sred_potreb_cen_08-2026.xlsx",
]
REFERENCE_FILES = [
    "spr_vid_tov_org_reg_cen_2026.xlsx", "spr_priobr_reg_sx_2025.xlsx",
    "spr_vid_tov_prod_cx_2026.xlsx", "spr_priobr_str-2026.xlsx",
    "harakteristiki_2023-2026.xlsx", "Nom-2026.xlsx",
    "s_nedw_funkc_nazn_2025.xlsx",
]
FUELS = {"7800": "All petrol", "7802": "AI-92", "7803": "AI-95",
         "7804": "Diesel", "7806": "AI-98 and above"}
CORE_FUELS = {"7802", "7803", "7804"}
ALIASES = {
    "городмосквастолицароссийскойфедерациигородфедеральногозначения": "гмосква",
    "городсанктпетербурггородфедеральногозначения": "гсанктпетербург",
    "республикаадыгеяадыгея": "республикаадыгея",
    "республикататарстантатарстан": "республикататарстан",
    "чувашскаяреспубликачувашия": "чувашскаяреспублика",
}


def norm(value):
    return re.sub(r"[^а-яa-z0-9]", "", str(value).lower().replace("ё", "е"))


def file_hash(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def code(value):
    if isinstance(value, (int, float)) and math.isfinite(value):
        return str(int(value))
    return str(value).strip() if value is not None else ""


def numeric(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value) if math.isfinite(value) else None
    return None


def scope(geo_code, geo_name):
    if geo_code == "643":
        return "national"
    if "федеральный округ" in geo_name.lower():
        return "federal_district"
    if geo_code.endswith("00000000") or "автономный округ" in geo_name.lower():
        return "region"
    return "surveyed_city"


def independent_raw_cell_checks(root, samples):
    """Check pilot sample cells through XLSX XML, independently of openpyxl."""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    selected = collections.defaultdict(list)
    for sample in samples:
        selected[(sample["source_file"], sample["sheet"])].append(sample)
    results = []
    for (filename, sheetname), records in selected.items():
        with zipfile.ZipFile(root / filename) as archive:
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            rid = next(s.attrib[f"{{{rel_ns}}}id"] for s in workbook.findall("m:sheets/m:sheet", ns)
                       if s.attrib["name"] == sheetname)
            rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(r.attrib["Target"] for r in rels if r.attrib["Id"] == rid)
            member = target.lstrip("/") if target.startswith("/") else "xl/" + target
            cells = {sample["cell"]: sample for sample in records}
            found = {}
            with archive.open(member) as stream:
                for _, cell in ET.iterparse(stream, events=("end",)):
                    if cell.tag != f"{{{ns['m']}}}c":
                        continue
                    if cell.attrib.get("r") in cells:
                        node = cell.find("m:v", ns)
                        value = float(node.text) if node is not None and node.text else None
                        found[cell.attrib["r"]] = value
                    cell.clear()
            for address, sample in cells.items():
                observed = found.get(address)
                expected = sample["extracted_value"]
                results.append({"source_file": filename, "sheet": sheetname, "cell": address,
                                "xml_numeric_value": observed, "extracted_value": expected,
                                "passed": observed == expected})
    return results


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    out = args.output or root / "outputs/new_prices_audit_2026-10-05"
    out.mkdir(parents=True, exist_ok=True)
    concordance = pd.read_csv(root / "outputs/regional_concordance_2026-09-29_v2/region_concordance.csv", dtype=str)
    name_map = {}
    eligible = set()
    for rec in concordance.to_dict("records"):
        rid = rec.get("region_id")
        if pd.isna(rid):
            continue
        accepted = rec.get("eligible_for_candidate_profile") == "True"
        if accepted:
            eligible.add(rid)
        for field in ("fns_name_ru", "rosstat_name_ru", "treasury_name_ru"):
            value = rec.get(field)
            if pd.notna(value):
                name_map[norm(value)] = (rid, accepted)
    fuel_rows, products, sheets, sources, problems, samples = [], [], [], [], [], []
    for filename in PRICE_FILES:
        path = root / filename
        digest = file_hash(path)
        sources.append({"filename": filename, "sha256": digest, "bytes": path.stat().st_size,
                        "role": "monthly_average_consumer_price_levels", "raw_preserved": True})
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook:
            period_match = re.fullmatch(r"(\d{2})\((\d{4})\)", sheet.title)
            if not period_match:
                continue
            month, year = map(int, period_match.groups())
            period = f"{year:04d}-{month:02d}-01"
            iterator = iter(sheet.iter_rows(values_only=True))
            prelude = []
            header = None
            header_row = None
            for row_number in range(1, 12):
                row = next(iterator)
                prelude.append(row)
                if row[0] == "Код территории" and row[1] == "Наименование территории":
                    header, header_row = row, row_number
                    break
            if header is None:
                raise ValueError(f"Cannot identify territory/product header: {filename}:{sheet.title}")
            labels = next(iterator)
            product_columns = {i: code(v) for i, v in enumerate(header) if i >= 2 and numeric(v) is not None}
            units_note = " ".join(str(v) for row in prelude for v in row if v and "рублей" in str(v))
            duplicate_products = [k for k, n in collections.Counter(product_columns.values()).items() if n > 1]
            for column, pcode in product_columns.items():
                products.append({"source_file": filename, "sheet": sheet.title, "period": period,
                                 "product_code": pcode, "product_name_ru": labels[column],
                                 "source_column": get_column_letter(column + 1)})
            counts = collections.Counter()
            geo_seen = set()
            missing_tokens = collections.Counter()
            regional_names = set()
            for row_number, row in enumerate(iterator, header_row + 2):
                geo_code = code(row[0])
                geo_name = str(row[1]).strip() if row[1] is not None else ""
                if not geo_code.isdigit() or not geo_name:
                    continue
                geography_scope = scope(geo_code, geo_name)
                if geography_scope == "region":
                    regional_names.add(geo_name)
                match = name_map.get(ALIASES.get(norm(geo_name), norm(geo_name)))
                rid, accepted = match if match else (None, False)
                if geo_code in geo_seen:
                    counts["duplicate_geography_rows"] += 1
                geo_seen.add(geo_code)
                counts[geography_scope + "_rows"] += 1
                for column, pcode in product_columns.items():
                    raw_value = row[column] if column < len(row) else None
                    value = numeric(raw_value)
                    counts["expected_price_cells"] += 1
                    if value is None:
                        counts["missing_or_nonnumeric_price_cells"] += 1
                        if raw_value is not None:
                            missing_tokens[str(raw_value)] += 1
                    elif value <= 0:
                        counts["nonpositive_price_cells"] += 1
                        problems.append({"source_file": filename, "sheet": sheet.title,
                                         "cell": f"{get_column_letter(column + 1)}{row_number}",
                                         "geography": geo_name, "product_code": pcode,
                                         "value": value, "issue": "nonpositive_price"})
                    else:
                        counts["positive_price_cells"] += 1
                    if pcode in FUELS:
                        fuel_rows.append({
                            "period": period, "source_geography_code": geo_code,
                            "geography_name_ru": geo_name, "geography_scope": geography_scope,
                            "region_id": rid if geography_scope == "region" else None,
                            "eligible_for_existing_dashboard": accepted if geography_scope == "region" else False,
                            "product_code": pcode, "fuel_grade": FUELS[pcode],
                            "product_name_ru": labels[column], "price_rub_per_litre": value,
                            "measurement_basis": "average_consumer_price_at_end_of_month",
                            "source_file": filename, "source_sha256": digest,
                            "source_sheet": sheet.title,
                            "source_cell": f"{get_column_letter(column + 1)}{row_number}",
                            "quality_flag": "reported" if value is not None and value > 0 else "missing_or_invalid",
                        })
                        if rid in {"RU-KLU", "RU-SVE"} and month in {1, 8, 12}:
                            samples.append({"source_file": filename, "sheet": sheet.title,
                                            "cell": f"{get_column_letter(column + 1)}{row_number}",
                                            "region_id": rid, "period": period, "product_code": pcode,
                                            "raw_value": raw_value, "extracted_value": value})
            sheets.append({"source_file": filename, "sheet": sheet.title, "period": period,
                           "header_row": header_row, "declared_max_rows": sheet.max_row,
                           "declared_max_columns": sheet.max_column,
                           "products": len(product_columns), "duplicate_product_codes": duplicate_products,
                           "units_note": units_note, "missing_tokens": dict(missing_tokens), **dict(counts)})
        workbook.close()
        print(f"Audited {filename}", flush=True)

    # Reference catalogues are never loaded as price observations.
    reference_roles = {
        "spr_vid_tov_org_reg_cen_2026.xlsx": "goods acquired by organisations catalogue",
        "spr_priobr_reg_sx_2025.xlsx": "agricultural organisations acquired goods and services catalogue",
        "spr_vid_tov_prod_cx_2026.xlsx": "agricultural producer products catalogue",
        "spr_priobr_str-2026.xlsx": "construction material specifications catalogue",
        "harakteristiki_2023-2026.xlsx": "consumer item specifications and CPI inclusion flags, not index observations",
        "Nom-2026.xlsx": "industrial producer product nomenclature, not producer price observations",
        "s_nedw_funkc_nazn_2025.xlsx": "commercial real estate function classification, not price observations",
    }
    for filename in REFERENCE_FILES:
        path = root / filename
        workbook = load_workbook(path, read_only=True, data_only=True)
        sources.append({"filename": filename, "sha256": file_hash(path), "bytes": path.stat().st_size,
                        "role": reference_roles[filename], "raw_preserved": True,
                        "sheets": [{"name": s.title, "rows": s.max_row, "columns": s.max_column} for s in workbook]})
        workbook.close()
    pdf = root / "cpi_ru(3).pdf"
    if pdf.exists():
        sources.append({"filename": pdf.name, "sha256": file_hash(pdf), "bytes": pdf.stat().st_size,
                        "role": "international CPI methodology manual, Russian edition 2007, 720 pages; not regional data",
                        "raw_preserved": True})

    frame = pd.DataFrame(fuel_rows)
    frame["period"] = pd.to_datetime(frame["period"])
    duplicate_mask = frame.duplicated(["source_geography_code", "period", "product_code"], keep=False)
    regional = frame[(frame.geography_scope == "region") & frame.eligible_for_existing_dashboard].copy()
    core = regional[regional.product_code.isin(CORE_FUELS)].copy()
    core_missing = core[core.price_rub_per_litre.isna()]
    previous = core[["region_id", "product_code", "period", "price_rub_per_litre"]].copy()
    previous["period"] = previous["period"] + pd.DateOffset(months=1)
    previous = previous.rename(columns={"price_rub_per_litre": "previous_month_price"})
    temporal = core.merge(previous, on=["region_id", "product_code", "period"],
                          how="left", validate="one_to_one")
    temporal["month_on_month_price_change_pct"] = (
        temporal.price_rub_per_litre / temporal.previous_month_price - 1
    ) * 100
    temporal_flags = temporal[temporal.month_on_month_price_change_pct.abs() > 10].copy()
    temporal_flags["review_reason"] = "absolute month-on-month average-price change exceeds 10%; retain pending interpretation"
    largest = temporal_flags.reindex(
        temporal_flags.month_on_month_price_change_pct.abs().sort_values(ascending=False).index
    ).head(24)
    for rec in largest.to_dict("records"):
        samples.append({"source_file": rec["source_file"], "sheet": rec["source_sheet"],
                        "cell": rec["source_cell"], "region_id": rec["region_id"],
                        "period": rec["period"].date().isoformat(), "product_code": rec["product_code"],
                        "raw_value": rec["price_rub_per_litre"], "extracted_value": rec["price_rub_per_litre"]})
    unmapped = frame[(frame.geography_scope == "region") & frame.region_id.isna()][["source_geography_code", "geography_name_ru"]].drop_duplicates()
    coverage = regional.groupby(["region_id", "product_code", "fuel_grade"], dropna=False).agg(
        expected_observations=("price_rub_per_litre", "size"),
        numeric_observations=("price_rub_per_litre", "count"),
        first_period=("period", "min"), last_period=("period", "max"),
        minimum=("price_rub_per_litre", "min"), maximum=("price_rub_per_litre", "max"),
    ).reset_index()
    product_frame = pd.DataFrame(products)
    changes = []
    for pcode, group in product_frame.groupby("product_code"):
        labels = group.product_name_ru.dropna().drop_duplicates().tolist()
        if len(labels) > 1:
            changes.append({"product_code": pcode, "different_labels": labels})
    independent = independent_raw_cell_checks(root, samples)
    hashes_unchanged = all(file_hash(root / rec["filename"]) == rec["sha256"] for rec in sources)
    summary = {
        "audit_executed_at_utc": datetime.now(timezone.utc).isoformat(),
        "status": "candidate_not_promoted", "source_files_unchanged": hashes_unchanged,
        "measure": "monthly average consumer price levels, not official CPI or PPI",
        "period_start": frame.period.min().date().isoformat(),
        "period_end": frame.period.max().date().isoformat(),
        "months": int(frame.period.nunique()), "fuel_rows_all_scopes": len(frame),
        "fuel_duplicate_grain_rows": int(duplicate_mask.sum()),
        "fuel_nonpositive_rows": int((frame.price_rub_per_litre <= 0).sum()),
        "fuel_missing_rows_all_scopes": int(frame.price_rub_per_litre.isna().sum()),
        "existing_dashboard_eligible_regions": len(eligible),
        "mapped_eligible_regions": int(regional.region_id.nunique()),
        "mapped_eligible_missing_regions": sorted(eligible - set(regional.region_id)),
        "eligible_core_fuel_observations_expected": len(core),
        "eligible_core_fuel_observations_missing": len(core_missing),
        "eligible_core_fuel_observations_nonpositive": int((core.price_rub_per_litre <= 0).sum()),
        "eligible_all_fuel_observations_expected": len(regional),
        "eligible_all_fuel_observations_missing": int(regional.price_rub_per_litre.isna().sum()),
        "all_product_expected_cells": sum(s["expected_price_cells"] for s in sheets),
        "all_product_missing_or_nonnumeric_cells": sum(s.get("missing_or_nonnumeric_price_cells", 0) for s in sheets),
        "all_product_nonpositive_cells": sum(s.get("nonpositive_price_cells", 0) for s in sheets),
        "sheets_with_duplicate_geography": sum(s.get("duplicate_geography_rows", 0) > 0 for s in sheets),
        "sheets_with_duplicate_product_codes": sum(bool(s["duplicate_product_codes"]) for s in sheets),
        "product_codes_with_label_drift": len(changes),
        "independent_raw_cell_checks": len(independent),
        "independent_raw_cell_checks_failed": sum(not c["passed"] for c in independent),
        "matched_month_on_month_changes": int(temporal.previous_month_price.notna().sum()),
        "month_on_month_review_flags_abs_above10pct": len(temporal_flags),
        "largest_month_on_month_change_pct": float(temporal.month_on_month_price_change_pct.max()),
        "smallest_month_on_month_change_pct": float(temporal.month_on_month_price_change_pct.min()),
        "notes": [
            "National, district, regional and surveyed-city rows are retained separately; do not average city rows into regional prices.",
            "Eligible regional subset follows existing 78-region dashboard concordance; no national aggregation claim.",
            "Missing or nonnumeric cells remain missing, never zero-filled or forward-filled.",
            "Prices are stated at end of period; latest workbook stops at August 2026, not September or October.",
            "Changing product specifications mean percentage changes in average prices are not official quality-adjusted CPI indices.",
            "Some city codes and non-fuel product columns change across vintages; joins use identifiers, not cell positions.",
            "No regional CPI time series was supplied in these files. Consumer specifications and CPI manual are not inflation observations.",
            "Large monthly changes are retained with analyst-review flags; a raw-source match verifies extraction, not an explanation for the change.",
        ],
    }
    try:
        frame.to_parquet(out / "fuel_prices_all_geographies_candidate.parquet", index=False)
        regional.to_parquet(out / "fuel_prices_eligible_regions_candidate.parquet", index=False)
        summary["candidate_storage_format"] = "parquet"
    except ImportError:
        # Bundled runtime can lack Arrow; retain typed columns via documented CSV
        # parsing until conversion in the existing project runtime is possible.
        frame.to_csv(out / "fuel_prices_all_geographies_candidate.csv.gz", index=False)
        regional.to_csv(out / "fuel_prices_eligible_regions_candidate.csv.gz", index=False)
        summary["candidate_storage_format"] = "gzip_csv"
    coverage.to_csv(out / "fuel_coverage.csv", index=False)
    core_missing.to_csv(out / "core_fuel_missing_cells.csv", index=False)
    temporal_flags.to_csv(out / "month_on_month_review_flags.csv", index=False)
    unmapped.to_csv(out / "unmapped_regional_geographies.csv", index=False)
    registry = product_frame.groupby(
        ["source_file", "product_code", "product_name_ru", "source_column"], dropna=False
    ).agg(first_period=("period", "min"), last_period=("period", "max"),
          months=("period", "nunique")).reset_index()
    registry.to_csv(out / "product_column_registry.csv", index=False)
    pd.DataFrame(samples).to_csv(out / "raw_cell_samples.csv", index=False)
    pd.DataFrame(problems).to_csv(out / "nonpositive_price_cells.csv", index=False)
    for name, value in [("summary", summary), ("source_manifest", sources), ("sheet_profiles", sheets),
                        ("product_label_drift", changes), ("independent_raw_cell_checks", independent)]:
        (out / f"{name}.json").write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
