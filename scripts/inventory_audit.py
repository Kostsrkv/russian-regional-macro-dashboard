"""Read-only source audit. Writes only its own dated audit outputs, never inputs."""
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import Counter, defaultdict
from io import BytesIO, StringIO
import csv
import hashlib
import json
import re
import sys
import subprocess
import zipfile
import os
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs' / os.environ.get('MACRO_AUDIT_OUTPUT', 'data_inventory_2026-09-26')
sys.path.insert(0, str(ROOT / 'src'))
# Avoid cloud-only bytecode caches; read the actual project source instead.
sys.pycache_prefix = '/private/tmp/macro-audit-bytecode-20260926'
sys.dont_write_bytecode = True
GRAINS={'pit_receipts':['region_id','reporting_cutoff','okved_raw','source_vintage'],
        'industrial_production':['region_id','period','okved_section','index_measure','source_vintage'],
        'budget_execution':['region_id','period','budget_level','revenue_code','source_vintage'],
        'regions':['region_id'],'sources':['source_id']}

def _pit_fields(schema):
    descriptions=schema['russian description'].fillna('')
    rows=schema[descriptions.str.contains('налог на доходы физических лиц',case=False)&descriptions.str.contains(r'\bвсего\b',case=False)&~descriptions.str.contains('в том числе',case=False)]
    sections={};unmapped=[]
    for _,r in rows.iterrows():
        desc=str(r['russian description'])
        if 'не распределен' in desc.lower() or 'не имеющ' in desc.lower():unmapped.append(r['field name']);continue
        m=re.match(r'^\s*([A-ZА-Я])\s+',desc)
        if m:sections.setdefault({'А':'A','В':'B','С':'C','Е':'E','Н':'H'}.get(m[1],m[1]),r['field name'])
    return sections,unmapped

def read_bounded(p):
    return subprocess.run([sys.executable,'-c','import sys; sys.stdout.buffer.write(open(sys.argv[1],"rb").read())',str(p)],capture_output=True,check=True,timeout=30).stdout

EXCLUDE = {'.git', '.venv', '__pycache__', '.pytest_cache', 'outputs'}
EXTS = {'.csv','.xlsx','.xls','.zip','.parquet','.json','.pdf','.docx','.crdownload','.html','.htm','.7z','.rar','.dta'}

def family(p):
    n=p.name
    if 'promoted' in p.parts: return 'promoted'
    if n.startswith(('data-','structure-','1nom')): return 'FNS'
    if n.startswith(('ind_','sezon_')): return 'Rosstat'
    if any(t in str(p) for t in ['КБРФ','КБСРФ','KBSRF']): return 'Treasury'
    return 'other'

def inspect_file(p):
    rec={'path':str(p.relative_to(ROOT)), 'family':family(p),'bytes':p.stat().st_size,
         'initially_dataless':bool(getattr(p.stat(),'st_flags',0)&0x40000000)}
    m=re.search(r'data-(\d{8})-structure-(\d{8}).*\.csv$',p.name)
    if m:
        rec['release_label']=m[1];rec['schema_label']=m[2]
        rec['matching_schema_present']=(ROOT/f'structure-{m[2]}.csv').exists()
    if rec['initially_dataless'] and '--probe-cloud' not in sys.argv:
        rec['read_status']='cloud_only_not_content_verified'
        return rec
    try:
        payload=read_bounded(p)
        rec['sha256']=hashlib.sha256(payload).hexdigest()
        rec['read_status']='readable'
        if p.suffix=='.csv':
            text=payload.decode('utf-8-sig')
            sep=';' if text.splitlines()[0].count(';')>text.splitlines()[0].count(',') else ','
            rows=list(csv.reader(StringIO(text),delimiter=sep)); header=rows[0]; body=rows[1:]
            rec.update(rows=len(body),columns=len(header),delimiter=sep,
                       irregular_width_rows=sum(len(r)!=len(header) for r in body),
                       duplicate_headers=len(header)-len(set(header)),
                       duplicate_rows=len(body)-len(set(map(tuple,body))))
            if 'GA' in header and 'GB' in header:
                codes=[r[header.index('GA')] for r in body if len(r)==len(header)]
                rec['source_region_rows']=len(set(codes));rec['duplicate_region_codes']=len(codes)-len(set(codes))
                nums=[v for r in body for j,v in enumerate(r) if j>1]
                rec['blank_measure_cells']=sum(v.strip()=='' for v in nums)
                rec['non_numeric_measure_cells']=sum(bool(v.strip()) and not re.fullmatch(r'-?\d+(?:[.,]\d+)?',v.strip()) for v in nums)
                rec['zero_measure_cells']=sum(v.strip() in ('0','0.0','0,0') for v in nums)
                m=re.search(r'data-(\d{8})-structure-(\d{8})',p.name)
                if m:
                    rec['release_label']=m[1];rec['schema_label']=m[2]
                    sp=ROOT/f'structure-{m[2]}.csv';rec['matching_schema_present']=sp.exists()
                    if sp.exists() and not (getattr(sp.stat(),'st_flags',0)&0x40000000):
                        schema=pd.read_csv(BytesIO(read_bounded(sp)),dtype=str)
                        fields=list(schema['field name'])
                        rec['schema_fields']=len(fields)
                        rec['header_matches_schema']=header==fields
                        sections,unmapped=_pit_fields(schema)
                        rec['pit_sections']=';'.join(sorted(sections));rec['pit_unmapped_fields']=len(unmapped)
            elif 'field name' in header:
                rec['schema_duplicate_fields']=len(body)-len({r[0] for r in body})
        elif p.suffix=='.xlsx':
            wb=openpyxl.load_workbook(BytesIO(payload),read_only=True,data_only=True)
            sheets=[]
            for ws in wb:
                preview=[]; region_hits=[]
                for i,row in enumerate(ws.iter_rows(values_only=True)):
                    if i<12: preview.append([str(v)[:220] if v is not None else None for v in row[:12]])
                    for v in row[:3]:
                        if isinstance(v,str) and any(s in v.lower() for s in ['калуж','свердлов','российская федерация']):
                            region_hits.append({'row':i+1,'label':v})
                sheets.append({'sheet':ws.title,'rows':ws.max_row,'columns':ws.max_column,'header_preview':preview,'region_hits':region_hits})
            rec['sheets']=sheets;wb.close()
        elif p.suffix=='.zip':
            def _decode_zip_name(n):
                try:return n.encode('cp437').decode('cp866')
                except UnicodeError:return n
            with zipfile.ZipFile(BytesIO(payload)) as z:
                rec['archive_members']=[_decode_zip_name(n) for n in z.namelist()]
                try: rec['outer_crc_error']=z.testzip()
                except Exception as e: rec['outer_crc_error']=str(e)
                if family(p)=='Treasury':
                    matches=[n for n in z.namelist() if '0503317' in _decode_zip_name(n) and 'html' in _decode_zip_name(n).lower()]
                    member=matches[0] if len(matches)==1 else None
                    if member:
                        with zipfile.ZipFile(BytesIO(z.read(member))) as nested:
                            rec['html_crc_error']=nested.testzip()
                            catalog_text=nested.read('TerrList_428m.html').decode('cp1251')
                            catalog={re.sub('<[^>]+>','',label).strip():re.split(r'[\\/]',url)[-1] for url,label in re.findall(r'<a[^>]+href=[\"\']([^\"\']+)[\"\'][^>]*>(.*?)</a>',catalog_text,re.I|re.S) if 'Mes_428m_' in url}
                            rec['territory_catalogue_entries']=len(catalog)
                            rec['territory_names']=list(catalog)
                            sample=[]
                            for name in ['Калужская область','Свердловская область']:
                                if name not in catalog:continue
                                html=nested.read(catalog[name]).decode('cp1251')
                                text=re.sub('<[^>]+>',' ',html)
                                patterns=['Расходы бюджета','Источники финансирования','Результат исполнения бюджета','ОКЕИ']
                                sample.append({'region':name,'section_markers':{t:t.lower() in text.lower() for t in patterns}})
                            rec['report_sections_sample']=sample
        elif p.suffix=='.xls':
            rec['format_signature']=payload[:8].hex()
            try:
                wb=pd.ExcelFile(BytesIO(payload));rec['sheet_names']=wb.sheet_names
            except Exception as e: rec['inspection_limit']=str(e)
        elif p.suffix=='.crdownload':
            rec['inspection_limit']='Unfinished browser download; excluded from analysis.'
        elif p.suffix=='.parquet':
            df=pd.read_parquet(BytesIO(payload));name=p.stem
            rec['rows']=len(df);rec['columns']=len(df.columns)
            profile={'table':name,'rows':len(df),'columns':len(df.columns),'null_counts':df.isna().sum().to_dict()}
            if name in GRAINS:profile['duplicate_grain_rows']=int(df.duplicated(GRAINS[name]).sum())
            for c in ['region_id','period','reporting_cutoff','revenue_code','okved_section','index_measure','quality_status']:
                if c in df:profile[c]=sorted(df[c].dropna().astype(str).unique().tolist())
            rec['profile']=profile
        elif p.name=='manifest.json':rec['manifest']=json.loads(payload)
    except Exception as e:
        rec['read_status']='error';rec['error']=repr(e)
    return rec

def run():
    OUT.mkdir(parents=True,exist_ok=True)
    paths=sorted(p for p in ROOT.rglob('*') if p.is_file() and not (set(p.relative_to(ROOT).parts)&EXCLUDE) and p.suffix.lower() in EXTS)
    records=[]
    with ThreadPoolExecutor(max_workers=8) as pool:
        tasks={pool.submit(inspect_file,p):p for p in paths}
        for i,f in enumerate(as_completed(tasks),1):
            rec=f.result();records.append(rec)
            if i%15==0: print(f'Inspected {i}/{len(paths)} files',flush=True)
    records.sort(key=lambda r:r['path'])
    (OUT/'file_details.json').write_text(json.dumps(records,ensure_ascii=False,indent=2,default=str))
    scalar=[{k:v for k,v in r.items() if not isinstance(v,(dict,list))} for r in records]
    pd.DataFrame(scalar).to_csv(OUT/'file_inventory.csv',index=False)
    hashes=defaultdict(list)
    for r in records:
        if r.get('sha256'): hashes[r['sha256']].append(r['path'])
    dup=[{'sha256':h,'paths':ps,'extra_copies':len(ps)-1} for h,ps in hashes.items() if len(ps)>1]
    (OUT/'duplicates.json').write_text(json.dumps(dup,ensure_ascii=False,indent=2))
    fns=[r for r in records if r.get('release_label')]
    releases={r['release_label']:r for r in fns}
    missing=sorted({r['schema_label'] for r in fns if not r['matching_schema_present']})
    checklist=[]
    for d in missing:
        checklist.append({'priority':'1' if d[4:]== '0101' or d[4:] in ['0401','0701','1001'] else '2',
          'source':'FNS','file':f'structure-{d}.csv','reason':'Required schema for existing data',
          'url_candidate_not_live_verified':f'https://data.nalog.ru/opendata/7707329152-1nomo/structure-{d}.csv'})
    for d in ['20260701']:
        for n in [f'data-{d}-structure-{d}.csv',f'structure-{d}.csv']:
            if not (ROOT/n).exists(): checklist.append({'priority':'1','source':'FNS','file':n,'reason':'Latest pair identified in prior source check','url_candidate_not_live_verified':f'https://data.nalog.ru/opendata/7707329152-1nomo/{n}'})
    # Calendar targets, not an assertion that FNS published every quarter.
    quarterly=[]
    for y in range(2017,2027):
        for md in ['0101','0401','0701','1001']:
            d=f'{y}{md}'
            if '20170401'<=d<='20260701':
                quarterly.append({'snapshot_label':d,'data_present':d in releases,'schema_present':(ROOT/f'structure-{d}.csv').exists(),
                                  'note':'Availability/publication schedule must be checked; filenames are not verified reporting cutoffs.'})
    pd.DataFrame(quarterly).to_csv(OUT/'fns_quarterly_coverage.csv',index=False)
    pd.DataFrame(checklist).sort_values(['priority','file']).to_csv(OUT/'download_checklist.csv',index=False)
    profiles=[r['profile'] for r in records if 'profile' in r]
    (OUT/'promoted_profiles.json').write_text(json.dumps(profiles,indent=2,ensure_ascii=False,default=str))
    # Do not let missing cloud payloads stall a metadata inventory or imply a fresh build.
    fresh_match=None
    summary={'scope':str(ROOT),'files':len(records),'families':dict(Counter(r['family'] for r in records)),
             'read_errors':[r['path'] for r in records if r['read_status']=='error'],
             'duplicate_groups':len(dup),'extra_identical_copies':sum(x['extra_copies'] for x in dup),
             'fns_data_csv_files':len(fns),'fns_distinct_releases':len(releases),'fns_missing_schema_count':len(missing),
             'fns_schema_backed_releases':sorted(d for d,r in releases.items() if r['matching_schema_present']),
             'cloud_only_files':sum(r['read_status']=='cloud_only_not_content_verified' for r in records),
             'readable_files':sum(r['read_status']=='readable' for r in records),
             'initially_dataless_files':sum(r['initially_dataless'] for r in records),
             'fresh_fns_matches_promoted_on_key_and_ytd':fresh_match,
             'raw_to_promoted_rebuild_status':'not_run_by_inventory; see separate reconciliation',
             'promoted_tables_profiled':len(profiles),'validation_status':'bounded profiles only; full pipeline not rerun'}
    (OUT/'summary.json').write_text(json.dumps(summary,indent=2,ensure_ascii=False))
    print(json.dumps(summary,indent=2,ensure_ascii=False),flush=True)
    return summary

if __name__=='__main__':run()
