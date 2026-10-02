"""Published January-to-month regional production indices, with all vintages.

These are same-period previous-year ratios. No monthly averaging, chaining of
output levels, or extraction of standalone quarters is supported here.
"""
from __future__ import annotations

import json
import re
import subprocess
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .ingest_rosstat import MONTHS, SHEETS as MONTHLY_SHEETS, _publication_date
from .provenance import sha256_file
from .regions import REGION_RECORDS

SHEETS = dict(zip(("2", "5", "8", "11", "14"), MONTHLY_SHEETS.values()))
GRAIN = ["region_id", "period", "okved_section"]
MEASURE = "same_period_previous_year_pct"
DATASET = "industrial_production_cumulative.parquet"
CONCORDANCE = "outputs/regional_concordance_2026-09-29_v2/region_concordance.csv"
INPUTS = ("ind_sub_2018_12-2025.xlsx", "ind_sub_2023_07-2026.xlsx")
SOURCE_URL = "https://rosstat.gov.ru/enterprise_industrial"
TITLE_PHRASES = {
    "TOTAL": "промышленного производства", "B": "добыча полезных ископаемых",
    "C": "обрабатывающие производства", "D": "обеспечение электрической энергией",
    "E": "водоснабжение",
}
REQUIRED_COLUMNS = set(GRAIN + [
    "window_start", "window_months", "frequency", "index_measure", "index_value",
    "growth_yoy_pct", "source_file", "source_sha256", "source_vintage", "source_sheet",
    "source_cell", "source_region_name", "raw_value", "raw_window_header",
    "raw_year_header", "availability", "index_base", "quality_status",
])
AUDIT_FILES = ("industrial_production_cumulative.csv", "vintage_observations.csv",
               "overlap_audit.csv", "source_basis_evidence.csv", "coverage.csv",
               "missing_observations.csv", "pilot_q1_handchecks.csv")
CLOUD_AUDIT_FILES = ("source_basis_evidence.csv", "coverage.csv",
                     "missing_observations.csv", "pilot_q1_handchecks.csv")


def normalize(value) -> str:
    """Normalize typography only: territorial qualifiers remain part of the key."""
    text = unicodedata.normalize("NFKC", str(value)).casefold().replace("ё", "е")
    return re.sub(r"[\s.\-–—]+", "", text)


def window_month(header) -> int:
    """Parse a published January-to-month label; retain footnotes in raw headers."""
    label = re.sub(r"\d+$", "", str(header).strip().casefold())
    label = re.sub(r"\s*[–—-]\s*", "-", label)
    if label == "январь":
        return 1
    if not label.startswith("январь-") or label[7:] not in MONTHS:
        raise ValueError(f"Not a January-to-month cumulative header: {header!r}")
    month = MONTHS[label[7:]]
    if month == 1:
        raise ValueError("Invalid repeated January window")
    return month


def _headers(rows, section) -> list[dict]:
    title, basis = str(rows[1][0]), str(rows[2][0]).strip().casefold()
    if TITLE_PHRASES[section] not in title.casefold():
        raise ValueError(f"Unexpected cumulative section title: {section}: {title}")
    if basis != "в % к соответствующему периоду предыдущего года":
        raise ValueError("Cumulative index basis is not same-period previous-year")
    years, headers, year = {}, [], None
    for col in range(1, len(rows[4])):
        year_header = rows[3][col]
        if year_header is not None:
            match = re.fullmatch(r"(20\d{2})\s+год\d*", str(year_header).strip())
            if not match:
                raise ValueError(f"Invalid cumulative year header: {year_header!r}")
            year = int(match.group(1))
            if year in years:
                raise ValueError("Repeated cumulative year block")
            years[year] = {"header": str(year_header), "months": []}
        window = rows[4][col]
        if year is None or window is None:
            raise ValueError("Missing cumulative calendar header")
        month = window_month(window)
        years[year]["months"].append(month)
        headers.append(dict(col=col, year=year, month=month,
                            raw_year_header=years[year]["header"], raw_window_header=str(window)))
    if not headers or list(years) != list(range(min(years), max(years) + 1)):
        raise ValueError("Non-contiguous cumulative year coverage")
    for y, block in years.items():
        if block["months"] != list(range(1, max(block["months"]) + 1)):
            raise ValueError("Cumulative windows do not start in January or have gaps")
        if y != max(years) and len(block["months"]) != 12:
            raise ValueError("Incomplete historical cumulative year")
    return headers


def _value(raw) -> tuple[float, str]:
    numeric = pd.to_numeric(raw, errors="coerce")
    if isinstance(raw, bool) or (pd.notna(numeric) and (not np.isfinite(numeric) or numeric < 0)):
        raise ValueError(f"Invalid cumulative production index: {raw!r}")
    if pd.notna(numeric):
        return float(numeric), "reported"
    return np.nan, "unavailable" if raw is None else "unavailable_source_marker"


def read_cumulative(path: Path, records: list[dict]) -> tuple[pd.DataFrame, dict]:
    """Read only the five published cumulative sheets, including every missing cell."""
    path = Path(path)
    wanted = {normalize(r["rosstat_name_ru"]): r for r in records}
    if len(wanted) != len(records) or len({r["region_id"] for r in records}) != len(records):
        raise ValueError("Ambiguous cumulative regional mapping")
    digest, output, evidence = sha256_file(path), [], []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not set(SHEETS).issubset(workbook.sheetnames) or "Содержание" not in workbook.sheetnames:
            raise ValueError("Cumulative workbook is missing required sheets")
        vintage = _publication_date(workbook)
        if pd.isna(vintage):
            raise ValueError("Cumulative publication/update date is unavailable")
        for sheet_name, (section, name_en, name_ru) in SHEETS.items():
            rows = list(workbook[sheet_name].values)
            headers = _headers(rows, section)
            footnotes = [str(r[0]) for r in rows[5:] if isinstance(r[0], str) and re.match(r"^\d", r[0])]
            stated_base = section == "TOTAL" and any("2023 базисного года" in n for n in footnotes)
            basis = "2023 GVA weights (explicit headline footnote)" if stated_base else "Weight base not stated in this sheet"
            evidence.append(dict(source_file=path.name, sha256=digest, source_vintage=vintage.strftime("%Y-%m-%d"),
                                 source_sheet=sheet_name, okved_section=section, title=rows[1][0],
                                 comparison_basis=rows[2][0], index_measure=MEASURE, index_base=basis,
                                 window_headers=[h["raw_window_header"] for h in headers],
                                 year_headers=list(dict.fromkeys(h["raw_year_header"] for h in headers)),
                                 periods=[str(pd.Timestamp(h["year"], h["month"], 1) + pd.offsets.MonthEnd(0))[:10] for h in headers],
                                 footnotes=footnotes))
            found = set()
            for row_number, row in enumerate(rows[5:], start=6):
                record = wanted.get(normalize(row[0]))
                if record is None:
                    continue
                if record["region_id"] in found:
                    raise ValueError(f"Duplicate regional cumulative row in sheet {sheet_name}")
                found.add(record["region_id"])
                for h in headers:
                    raw = row[h["col"]]
                    numeric, availability = _value(raw)
                    period = pd.Timestamp(h["year"], h["month"], 1) + pd.offsets.MonthEnd(0)
                    output.append(dict(
                        region_id=record["region_id"], region_name_ru=record["treasury_name_ru"],
                        period=period, frequency="YTD", window_start=pd.Timestamp(h["year"], 1, 1),
                        window_months=h["month"], okved_section=section,
                        industry_name_en=name_en, industry_name_ru=name_ru,
                        index_value=numeric, index_measure=MEASURE, index_base=basis,
                        growth_yoy_pct=numeric - 100 if pd.notna(numeric) else np.nan,
                        source_vintage=vintage.strftime("%Y-%m-%d"), source_file=path.name,
                        source_sha256=digest, source_id=f"rosstat-cumulative-{digest[:12]}",
                        source_sheet=sheet_name, source_cell=f"{get_column_letter(h['col'] + 1)}{row_number}",
                        source_region_name=str(row[0]), region_match_method=record.get("rosstat_match_method", "typography_normalization"),
                        scope_qualifier=str(row[0]) if re.search(r"без|в том числе", str(row[0]).casefold()) else "",
                        raw_value="" if raw is None else str(raw),
                        raw_year_header=h["raw_year_header"], raw_window_header=h["raw_window_header"],
                        availability=availability, missing_source_value=availability != "reported",
                        year_header_has_footnote=bool(re.search(r"год\d+$", h["raw_year_header"])),
                        window_header_has_footnote=bool(re.search(r"\d+$", h["raw_window_header"])),
                        revision_state="published_source_vintage", quality_status="candidate",
                        source_url=SOURCE_URL))
            if found != {r["region_id"] for r in records}:
                raise ValueError(f"Cumulative sheet {sheet_name}: unmatched eligible regions")
    finally:
        workbook.close()
    result = pd.DataFrame(output)
    if result.empty or result.duplicated(GRAIN).any():
        raise ValueError("Empty or duplicate cumulative observations")
    return result, dict(source_file=path.name, sha256=digest, source_vintage=vintage.strftime("%Y-%m-%d"),
                        source_url=SOURCE_URL, retrieval_date=None, rows=len(result), sheet_evidence=evidence)


def select_vintages(frames: list[pd.DataFrame]) -> pd.DataFrame:
    """Newest dated vintage wins, including its missing or suppressed values."""
    data = pd.concat(frames, ignore_index=True)
    dates = pd.to_datetime(data.source_vintage, errors="raise")
    if dates.isna().any() or data.duplicated(GRAIN + ["source_vintage"]).any():
        raise ValueError("Ambiguous same-date cumulative vintages")
    return (data.assign(_vintage=dates).sort_values("_vintage")
            .drop_duplicates(GRAIN, keep="last").drop(columns="_vintage")
            .sort_values(GRAIN).reset_index(drop=True))


def overlap_audit(frames: list[pd.DataFrame]) -> pd.DataFrame:
    """Every overlapping original observation and its newest-vintage selection."""
    original = pd.concat(frames, ignore_index=True)
    selected = select_vintages(frames)
    overlap = original.loc[original.duplicated(GRAIN, keep=False)].copy()
    chosen = selected[GRAIN + ["index_value", "availability", "source_vintage", "source_file", "source_cell", "source_sha256"]]
    chosen = chosen.rename(columns={c: "selected_" + c for c in chosen.columns if c not in GRAIN})
    audit = overlap.merge(chosen, on=GRAIN, validate="many_to_one")
    audit["is_selected_vintage"] = audit.source_vintage.eq(audit.selected_source_vintage)
    audit["index_difference_from_selected"] = audit.index_value - audit.selected_index_value
    both = audit.index_value.notna() & audit.selected_index_value.notna()
    audit["comparison_status"] = np.select(
        [both & audit.index_difference_from_selected.eq(0), both,
         audit.index_value.notna() & audit.selected_index_value.isna(),
         audit.index_value.isna() & audit.selected_index_value.notna()],
        ["unchanged", "revised", "latest_missing", "original_missing"], default="both_missing")
    return audit.sort_values(GRAIN + ["source_vintage"]).reset_index(drop=True)


def _validate(data: pd.DataFrame, manifest: dict) -> None:
    if not REQUIRED_COLUMNS.issubset(data.columns):
        raise ValueError("Missing cumulative candidate columns")
    if len(data) != manifest["rows"] or data.duplicated(GRAIN).any() or data[GRAIN].isna().any().any():
        raise ValueError("Invalid cumulative candidate grain")
    if not data.frequency.eq("YTD").all() or not data.index_measure.eq(MEASURE).all():
        raise ValueError("Invalid cumulative candidate index basis")
    data["period"] = pd.to_datetime(data.period, errors="raise")
    data["window_start"] = pd.to_datetime(data.window_start, errors="raise")
    if (data.period.isna().any() or data.window_start.isna().any()
            or not data.period.dt.is_month_end.all()
            or not data.window_start.dt.month.eq(1).all() or not data.window_start.dt.day.eq(1).all()
            or not data.window_start.dt.year.eq(data.period.dt.year).all()
            or not data.window_months.eq(data.period.dt.month).all()):
        raise ValueError("Invalid cumulative observation windows")
    if not data.raw_window_header.map(window_month).eq(data.window_months).all():
        raise ValueError("Cumulative raw headers disagree with window dates")
    header_year = data.raw_year_header.str.extract(r"^(20\d{2})\s+год\d*$", expand=False)
    if header_year.isna().any() or not header_year.astype(int).eq(data.period.dt.year).all():
        raise ValueError("Cumulative raw year headers disagree with period dates")
    numeric = data.index_value.dropna()
    if not np.isfinite(numeric).all() or numeric.lt(0).any() or not np.allclose(data.growth_yoy_pct, data.index_value - 100, equal_nan=True):
        raise ValueError("Invalid cumulative values or growth calculation")
    if (not data.availability.isin(["reported", "unavailable", "unavailable_source_marker"]).all()
            or not data.availability.eq("reported").eq(data.index_value.notna()).all()
            or not data.quality_status.eq("candidate").all()):
        raise ValueError("Invalid cumulative availability or candidate status")
    raw_numeric = pd.to_numeric(data.raw_value, errors="coerce")
    if (not np.allclose(raw_numeric, data.index_value, equal_nan=True)
            or not data.availability.eq("unavailable").eq(data.raw_value.eq("")).all()):
        raise ValueError("Cumulative raw values or missing markers disagree")
    sources = manifest["sources"]
    if len({s["source_file"] for s in sources}) != len(sources):
        raise ValueError("Duplicate cumulative source manifests")
    source_hash = {s["source_file"]: s["sha256"] for s in sources}
    source_date = {s["source_file"]: s["source_vintage"] for s in sources}
    bases = {(s["source_file"], e["source_sheet"]): e["index_base"]
             for s in sources for e in s["sheet_evidence"]}
    if (not data.source_file.map(source_hash).eq(data.source_sha256).all()
            or not data.source_file.map(source_date).eq(data.source_vintage).all()
            or not data.source_cell.str.fullmatch(r"[A-Z]+[1-9]\d*").all()
            or not data.source_sheet.map({k: v[0] for k, v in SHEETS.items()}).eq(data.okved_section).all()):
        raise ValueError("Cumulative source provenance disagrees")
    expected_bases = pd.Series([bases.get((file, sheet)) for file, sheet in zip(data.source_file, data.source_sheet)], index=data.index)
    if not data.index_base.eq(expected_bases).all():
        raise ValueError("Cumulative source weight-base caveats disagree")
    ids = manifest["eligible_region_ids"]
    if len(ids) != len(set(ids)) or set(data.region_id) != set(ids):
        raise ValueError("Cumulative eligible region identifiers disagree")
    # The exact source calendar defines coverage; do not accept dropped missing cells.
    expected = {}
    for s in sources:
        for e in s["sheet_evidence"]:
            if e["index_measure"] != MEASURE or e["comparison_basis"].strip().casefold() != "в % к соответствующему периоду предыдущего года":
                raise ValueError("Unverified cumulative source basis")
            for period in e["periods"]:
                key = (period, e["okved_section"])
                if key not in expected or s["source_vintage"] > expected[key][0]:
                    expected[key] = (s["source_vintage"], s["source_file"])
    if len(data) != len(ids) * len(expected):
        raise ValueError("Incomplete cumulative source-calendar coverage")
    for (period, section), (vintage, source) in expected.items():
        group = data.loc[data.period.eq(pd.Timestamp(period)) & data.okved_section.eq(section)]
        if len(group) != len(ids) or not group.source_vintage.eq(vintage).all() or not group.source_file.eq(source).all():
            raise ValueError("Cumulative newest-vintage selection or coverage disagrees")


def load_cumulative_candidate(root: Path) -> tuple[pd.DataFrame, dict]:
    """Fail closed on candidate status, file integrity, basis, grain and selection."""
    root = Path(root)
    try:
        manifest = json.loads((root / "manifest.json").read_text())
        if manifest.get("status") != "candidate_not_promoted" or manifest.get("blocking_failures") != 0:
            raise ValueError("Cumulative candidate is not an unblocked review candidate")
        if sha256_file(root / DATASET) != manifest.get("dataset_sha256"):
            raise ValueError("Cumulative dataset integrity check failed")
        hashes = manifest["artifact_sha256"]
        mode = manifest.get("distribution_mode", "private_full_audit")
        if mode == "cloud_compact":
            audit_files = CLOUD_AUDIT_FILES
            if set(hashes) != set(audit_files) or not set(AUDIT_FILES).issubset(manifest.get("private_audit_artifact_sha256", {})):
                raise ValueError("Incomplete compact cloud audit lineage")
        elif mode == "private_full_audit":
            audit_files = AUDIT_FILES
        else:
            raise ValueError("Unknown cumulative distribution mode")
        if not set(audit_files).issubset(hashes) or any(sha256_file(root / f) != hashes[f] for f in audit_files):
            raise ValueError("Cumulative audit artifact integrity check failed")
        data = pd.read_parquet(root / DATASET)
        _validate(data, manifest)
        return data, manifest
    except (KeyError, TypeError) as exc:
        raise ValueError("Incomplete cumulative candidate manifest") from exc


def _handchecks(workspace: Path, data: pd.DataFrame) -> pd.DataFrame:
    """Independent direct-cell reread of latest Q1 2024–26 for both pilots."""
    sample = data.loc[data.region_id.isin(r["region_id"] for r in REGION_RECORDS)
                      & data.period.dt.month.eq(3) & data.period.dt.year.isin([2024, 2025, 2026])]
    result = []
    for file, group in sample.groupby("source_file"):
        workbook = load_workbook(workspace / file, read_only=True, data_only=True)
        try:
            for row in group.itertuples():
                sheet = workbook[row.source_sheet]
                # Region and calendar labels establish the cell's meaning, separately from parser output.
                raw = sheet[row.source_cell].value
                excel_row = int(re.search(r"\d+$", row.source_cell).group())
                col = re.match(r"[A-Z]+", row.source_cell).group()
                region = sheet[f"A{excel_row}"].value
                header = sheet[f"{col}5"].value
                passed = (normalize(region) == normalize(row.source_region_name)
                          and str(header).startswith("январь-март")
                          and pd.notna(raw) and float(raw) == row.index_value)
                result.append(dict(region_id=row.region_id, period=row.period,
                                   window_start=row.window_start, okved_section=row.okved_section,
                                   source_file=file, source_sha256=row.source_sha256,
                                   source_vintage=row.source_vintage, source_sheet=row.source_sheet,
                                   source_cell=row.source_cell, raw_region_label=region,
                                   raw_window_header=header, raw_index_value=raw,
                                   candidate_index_value=row.index_value, passed=passed))
        finally:
            workbook.close()
    checks = pd.DataFrame(result)
    if len(checks) != 30 or not checks.passed.all():
        raise ValueError("Pilot Q1 independent raw-cell checks failed")
    return checks


def _write_parquet(data: pd.DataFrame, path: Path, workspace: Path) -> None:
    """Keep workbook reads in the bundled runtime; use the project codec if absent."""
    try:
        data.to_parquet(path, index=False)
    except ImportError:
        # The bundled spreadsheet runtime does not ship a Parquet codec. The
        # existing project environment serializes reviewed JSON only: it does
        # not extract, reinterpret or write to the original workbooks.
        codec = workspace / ".venv/bin/python"
        if not codec.is_file():
            raise ImportError("Parquet export requires the existing project .venv codec")
        code = ("import sys,pandas as pd; from io import StringIO; "
                "data=pd.read_json(StringIO(sys.stdin.read()),orient='table'); "
                "data.to_parquet(sys.argv[1],index=False)")
        subprocess.run([str(codec), "-c", code, str(path)],
                       input=data.to_json(orient="table", date_format="iso", double_precision=15),
                       text=True, capture_output=True, check=True)


def build_cumulative_candidate(workspace: Path, output: Path) -> dict:
    """Build a fresh local evidence candidate; never promote or overwrite."""
    workspace, output = Path(workspace), Path(output)
    if output.exists():
        raise FileExistsError("Use a fresh cumulative candidate directory")
    crosswalk = pd.read_csv(workspace / CONCORDANCE, dtype={"fns_code": str}).fillna("")
    records = crosswalk.loc[crosswalk.eligible_for_candidate_profile.eq(True)].to_dict("records")
    if len(records) != 78 or any(r["eligible_for_national_aggregation"] for r in records):
        raise ValueError("Unexpected cumulative concordance eligibility")
    pilots = {r["region_name_ru"]: r["region_id"] for r in REGION_RECORDS}
    if any(r["region_id"] != pilots[r["treasury_name_ru"]] for r in records if r["treasury_name_ru"] in pilots):
        raise ValueError("Cumulative concordance changed pilot identifiers")
    frames, sources = [], []
    for filename in INPUTS:
        frame, source = read_cumulative(workspace / filename, records)
        frames.append(frame)
        sources.append(source)
    data = select_vintages(frames)
    original = pd.concat(frames, ignore_index=True).sort_values(GRAIN + ["source_vintage"])
    overlaps = overlap_audit(frames)
    checks = _handchecks(workspace, data)
    coverage = data.groupby(["region_id", "okved_section"], as_index=False).agg(
        observations=("period", "size"), reported=("index_value", "count"),
        earliest=("period", "min"), latest=("period", "max"))
    coverage["missing"] = coverage.observations - coverage.reported
    basis = pd.DataFrame([e for s in sources for e in s["sheet_evidence"]])
    for c in ["window_headers", "year_headers", "periods", "footnotes"]:
        basis[c] = basis[c].map(lambda v: json.dumps(v, ensure_ascii=False))
    manifest = dict(
        status="candidate_not_promoted", blocking_failures=0, rows=len(data),
        regions=data.region_id.nunique(), periods=data.period.nunique(),
        eligible_region_ids=sorted(data.region_id.unique()), sections=[v[0] for v in SHEETS.values()],
        frequency="YTD", index_measure=MEASURE, unit="percent index; previous-year same-window = 100",
        earliest=str(data.period.min())[:10], latest=str(data.period.max())[:10],
        original_vintage_rows=len(original), overlap_rows=len(overlaps),
        overlap_keys=len(overlaps[GRAIN].drop_duplicates()),
        reported_rows=int(data.index_value.notna().sum()), missing_rows=int(data.index_value.isna().sum()),
        sources=sources, concordance=dict(file=CONCORDANCE, sha256=sha256_file(workspace / CONCORDANCE)),
        pilot_q1_handchecks=len(checks), pilot_identifiers=pilots,
        vintage_selection="Newest dated workbook wins per region/period/section, including missing values",
        limitations=[
            "Published Jan-to-month ratios only. No monthly averaging, output-level chaining or standalone-quarter derivation.",
            "March represents Jan–Mar/Q1. June represents Jan–Jun/H1, September Jan–Sep, December the whole year.",
            "Latest TOTAL footnote explicitly states 2023 GVA weights. Earlier TOTAL and section B–E sheets do not state a weight base; comparable constant-method output levels are not established.",
            "All original source vintages and missing/suppressed cells remain in the evidence files; dated revision notes are retained without inventing a per-cell revision history.",
            "Concordance eligibility is regional-profile use only. Held nested and separate geographies are excluded. No national aggregation is supported.",
            "Official source URL is retained; local workbook download/retrieval time is unknown. Workbook update dates are source vintages, not retrieval dates.",
            "Only TOTAL and broad OKVED2 B–E sections are available. There are no detailed production-type output weights or source quantities.",
            "Candidate is staged locally for review; no promotion or deployment.",
        ])
    _validate(data, manifest)
    output.mkdir(parents=True, exist_ok=False)
    _write_parquet(data, output / DATASET, workspace)
    tables = {"industrial_production_cumulative.csv": data, "vintage_observations.csv": original,
              "overlap_audit.csv": overlaps, "source_basis_evidence.csv": basis,
              "coverage.csv": coverage, "missing_observations.csv": data.loc[data.index_value.isna()],
              "pilot_q1_handchecks.csv": checks}
    for filename, frame in tables.items():
        frame.to_csv(output / filename, index=False)
    manifest["dataset_sha256"] = sha256_file(output / DATASET)
    manifest["artifact_sha256"] = {f: sha256_file(output / f) for f in tables}
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    try:
        load_cumulative_candidate(output)
    except ImportError:
        # Verify the normal public loader in the existing Parquet-capable
        # project environment, after all extraction and checks have finished.
        code = ("import sys; from pathlib import Path; sys.path.insert(0,sys.argv[1]); "
                "from macro_rus.cumulative_production import load_cumulative_candidate; "
                "load_cumulative_candidate(Path(sys.argv[2]))")
        subprocess.run([str(workspace / ".venv/bin/python"), "-c", code,
                        str(workspace / "src"), str(output)], capture_output=True, text=True, check=True)
    return manifest
