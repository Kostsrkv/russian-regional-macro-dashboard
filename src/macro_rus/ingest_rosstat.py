"""Ingest Rosstat monthly regional industrial-production workbooks."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook

from .contracts import INDUSTRIAL_PRODUCTION_COLUMNS, SOURCE_COLUMNS
from .provenance import sha256_file
from .regions import REGION_RECORDS


SHEETS = {
    "1": ("TOTAL", "Industrial production", "Промышленное производство"),
    "4": ("B", "Mining and quarrying", "Добыча полезных ископаемых"),
    "7": ("C", "Manufacturing", "Обрабатывающие производства"),
    "10": ("D", "Electricity, gas and steam", "Обеспечение электроэнергией, газом и паром"),
    "13": ("E", "Water, sewerage and waste", "Водоснабжение, водоотведение и обращение с отходами"),
}

MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4,
    "май": 5, "июнь": 6, "июль": 7, "август": 8,
    "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
}


@dataclass(frozen=True)
class RosstatIngestionResult:
    industrial_production: pd.DataFrame
    source_metadata: pd.DataFrame
    skipped_files: tuple[str, ...]


def _selected_regions(regions: Iterable[str] | None, region_records=None) -> list[dict[str, object]]:
    records = list(REGION_RECORDS if region_records is None else region_records)
    requested = {str(value).casefold() for value in regions or ()}
    if not requested:
        return records
    return [
        record for record in records
        if requested & {
            str(record["region_id"]).casefold(),
            str(record["region_name_en"]).casefold(),
            str(record["region_name_ru"]).casefold(),
        }
    ]


def _publication_date(workbook) -> pd.Timestamp:
    sheet = workbook["Содержание"]
    text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    match = re.search(r"Обновлено:\s*(\d{2})\.(\d{2})\.(\d{4})", text)
    if not match:
        return pd.NaT
    return pd.Timestamp(year=int(match.group(3)), month=int(match.group(2)), day=int(match.group(1)))


def _sheet_rows(workbook, sheet_name: str, records: list[dict[str, object]]) -> list[dict[str, object]]:
    sheet = workbook[sheet_name]
    years: list[int | None] = []
    current_year: int | None = None
    for cell in sheet[4][1:]:
        match = re.search(r"(20\d{2})", str(cell.value or ""))
        if match:
            current_year = int(match.group(1))
        years.append(current_year)
    months = [MONTHS.get(str(cell.value or "").strip().casefold()) for cell in sheet[5][1:]]
    normalize = lambda value: re.sub(r"[\s.\-–—]+", "", str(value).casefold().replace("ё", "е"))
    wanted = {normalize(record.get("rosstat_name_ru") or record["region_name_ru"]): record for record in records}
    if len(wanted) != len(records):
        raise ValueError("Duplicate Rosstat names in regional mapping")
    section, name_en, name_ru = SHEETS[sheet_name]
    output: list[dict[str, object]] = []
    for row in sheet.iter_rows(min_row=6, values_only=True):
        region_name = normalize(row[0] or "")
        record = wanted.get(region_name)
        if record is None:
            continue
        for year, month, value in zip(years, months, row[1:]):
            numeric = pd.to_numeric(value, errors="coerce")
            if year is None or month is None or pd.isna(numeric):
                continue
            period = pd.Timestamp(year=year, month=month, day=1) + pd.offsets.MonthEnd(0)
            output.append(
                {
                    "region_id": record["region_id"],
                    "region_name_en": record["region_name_en"],
                    "region_name_ru": record["region_name_ru"],
                    "period": period,
                    "frequency": "monthly",
                    "okved_section": section,
                    "industry_name_en": name_en,
                    "industry_name_ru": name_ru,
                    "index_measure": "same_month_previous_year_pct",
                    "index_value": float(numeric),
                    "index_base": "2023-base basket" if year >= 2026 else "2018-base basket",
                }
            )
    return output


def ingest_rosstat_result(
    raw_dir: str | Path,
    regions: Iterable[str] | None = None,
    *,
    region_records=None,
) -> RosstatIngestionResult:
    root = Path(raw_dir)
    records = _selected_regions(regions, region_records)
    frames: list[pd.DataFrame] = []
    sources: list[dict[str, object]] = []
    skipped: list[str] = []
    seen_hashes: set[str] = set()

    for path in sorted(root.glob("ind_sub_*.xlsx")):
        digest = sha256_file(path)
        if digest in seen_hashes:
            skipped.append(f"{path.name}: duplicate content")
            continue
        seen_hashes.add(digest)
        workbook = load_workbook(path, read_only=True, data_only=True)
        if not set(SHEETS).issubset(workbook.sheetnames) or len(workbook.sheetnames) < 15:
            skipped.append(f"{path.name}: annual/non-monthly workbook")
            workbook.close()
            continue
        publication_date = _publication_date(workbook)
        source_date = publication_date.strftime("%Y%m%d") if pd.notna(publication_date) else "undated"
        source_id = f"rosstat_ip_{source_date}_{digest[:12]}"
        rows: list[dict[str, object]] = []
        for sheet_name in SHEETS:
            rows.extend(_sheet_rows(workbook, sheet_name, records))
        workbook.close()
        if not rows:
            skipped.append(f"{path.name}: pilot regions not found")
            continue
        frame = pd.DataFrame(rows)
        frame["source_id"] = source_id
        frame["source_vintage"] = f"rosstat-ip-{source_date}-{digest[:12]}"
        frame["revision_status"] = "latest"
        frame["quality_status"] = "verified"
        frame["is_official"] = True
        frames.append(frame.loc[:, list(INDUSTRIAL_PRODUCTION_COLUMNS)])
        sources.append(
            {
                "source_id": source_id,
                "publisher": "Federal State Statistics Service (Rosstat)",
                "dataset_name": "Regional industrial production indices by economic activity",
                "source_url": "https://rosstat.gov.ru/enterprise_industrial",
                "local_path": str(path.resolve()),
                "retrieved_at_utc": pd.NaT,
                "reporting_cutoff": frame["period"].max(),
                "publication_date": publication_date,
                "units": "% of same month in previous year",
                "sha256": digest,
                "file_size_bytes": path.stat().st_size,
                "is_official": True,
                "notes": "Published monthly index; source-native values retained.",
            }
        )

    combined = (
        pd.concat(frames, ignore_index=True)
        if frames else pd.DataFrame(columns=INDUSTRIAL_PRODUCTION_COLUMNS)
    )
    if not combined.empty:
        combined["_publication"] = combined["source_vintage"].str.extract(r"(\d{8})", expand=False)
        grain = ["region_id", "period", "okved_section", "index_measure"]
        combined = (
            combined.sort_values([*grain, "_publication", "source_vintage"], kind="stable")
            .drop_duplicates(grain, keep="last")
            .drop(columns="_publication")
            .sort_values(["region_id", "period", "okved_section"], kind="stable")
            .reset_index(drop=True)
        )
    source_frame = pd.DataFrame(sources, columns=SOURCE_COLUMNS)
    if not combined.empty and not source_frame.empty:
        source_frame = source_frame[source_frame["source_id"].isin(combined["source_id"].unique())].reset_index(drop=True)
    return RosstatIngestionResult(combined, source_frame, tuple(skipped))


def ingest_rosstat(raw_dir: str | Path, regions: Iterable[str] | None = None) -> pd.DataFrame:
    return ingest_rosstat_result(raw_dir, regions).industrial_production
