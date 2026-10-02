"""Source-native annual production audit; never average monthly growth rates."""
from __future__ import annotations

import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .ingest_rosstat import _publication_date, SHEETS as MONTHLY_SHEETS
from .provenance import sha256_file

SHEETS = dict(zip([str(n) for n in range(1, 6)], MONTHLY_SHEETS.values()))
GRAIN = ["region_id", "period", "okved_section"]
ANNUAL_NAME_ALIASES = {
    "Чукотский авт.округ": "Чукотский автономный округ",
    "Еврейская авт.область": "Еврейская автономная область",
    "г. Санкт-Петербург": "Город Санкт-Петербург",
    "г. Москва": "Город Москва",
}


def normalize(value):
    return re.sub(r"[\s.\-–—]+", "", str(value).casefold().replace("ё", "е"))


def read_annual(path: Path, records: list[dict]) -> tuple[pd.DataFrame, dict]:
    wanted = {normalize(ANNUAL_NAME_ALIASES.get(r["rosstat_name_ru"], r["rosstat_name_ru"])): r for r in records}
    if len(wanted) != len(records):
        raise ValueError("Duplicate annual Rosstat region names.")
    digest = sha256_file(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    output, notes = [], {}
    try:
        if not set(SHEETS).issubset(workbook.sheetnames):
            raise ValueError("Annual workbook is missing section sheets.")
        vintage = _publication_date(workbook)
        if pd.isna(vintage):
            raise ValueError("Annual publication/update date is unavailable.")
        for sheet_name, (section, name_en, name_ru) in SHEETS.items():
            rows = list(workbook[sheet_name].values)
            if "к предыдущему году" not in str(rows[2][0]).casefold():
                raise ValueError("Annual index basis is not previous year.")
            years = [(i, int(m.group(1)), str(v)) for i, v in enumerate(rows[3]) if (m := re.match(r"^(20\d{2})(?:\D|\d|$)", str(v))) and i > 0]
            if not years or len({y for _, y, _ in years}) != len(years):
                raise ValueError("Invalid annual year headers.")
            footnotes = [str(r[0]) for r in rows[4:] if isinstance(r[0], str) and re.match(r"^\d\s", r[0])]
            notes[sheet_name] = footnotes
            basis = "2023 GVA weights (explicit headline footnote)" if section == "TOTAL" and any("2023 базисного" in n for n in footnotes) else "Weight base not stated in this sheet"
            found = set()
            for row_number, row in enumerate(rows[4:], start=5):
                record = wanted.get(normalize(row[0]))
                if record is None:
                    continue
                if record["region_id"] in found:
                    raise ValueError("Duplicate regional annual row.")
                found.add(record["region_id"])
                for col, year, header in years:
                    raw = row[col]
                    numeric = pd.to_numeric(raw, errors="coerce")
                    if pd.notna(numeric) and (not np.isfinite(numeric) or numeric < 0):
                        raise ValueError("Invalid annual production index.")
                    missing = "unavailable_source_marker" if pd.isna(numeric) and raw is not None else "unavailable" if pd.isna(numeric) else "reported"
                    output.append(dict(region_id=record["region_id"], region_name_ru=record["treasury_name_ru"],
                        period=pd.Timestamp(year, 12, 31), frequency="annual", okved_section=section,
                        industry_name_en=name_en, industry_name_ru=name_ru, index_value=numeric,
                        index_measure="same_year_previous_year_pct", index_base=basis,
                        growth_yoy_pct=numeric - 100 if pd.notna(numeric) else np.nan,
                        source_vintage=vintage.strftime("%Y-%m-%d"), source_file=path.name, source_sha256=digest,
                        source_id=f"rosstat-annual-{digest[:12]}", source_sheet=sheet_name,
                        source_cell=f"{get_column_letter(col+1)}{row_number}", source_region_name=row[0],
                        region_match_method="explicit_annual_alias" if record["rosstat_name_ru"] in ANNUAL_NAME_ALIASES else "typography_normalization",
                        raw_value="" if raw is None else str(raw), raw_year_header=header,
                        availability=missing, revision_state="source_first_estimate" if year == 2025 and any("Первая оценка" in n for n in footnotes) else "published",
                        quality_status="candidate", source_url="https://rosstat.gov.ru/enterprise_industrial"))
            if len(found) != len(records):
                raise ValueError(f"Annual sheet {sheet_name}: unmatched eligible regions.")
    finally:
        workbook.close()
    result = pd.DataFrame(output)
    if result.duplicated(GRAIN).any():
        raise ValueError("Duplicate annual observations.")
    return result, dict(source_file=path.name, sha256=digest, source_vintage=vintage.strftime("%Y-%m-%d"),
        source_url="https://rosstat.gov.ru/enterprise_industrial", retrieval_date=None, footnotes=notes)


def select_vintages(frames):
    data = pd.concat(frames, ignore_index=True)
    if data.duplicated(GRAIN + ["source_vintage"]).any():
        raise ValueError("Ambiguous same-date annual vintages.")
    return data.sort_values("source_vintage").drop_duplicates(GRAIN, keep="last").sort_values(GRAIN).reset_index(drop=True)


def load_annual_candidate(root: Path):
    manifest = json.loads((root / "manifest.json").read_text())
    path = root / "industrial_production_annual.parquet"
    if manifest.get("status") != "candidate_not_promoted" or manifest.get("blocking_failures") != 0 or sha256_file(path) != manifest.get("dataset_sha256"):
        raise ValueError("Annual candidate manifest/integrity check failed.")
    data = pd.read_parquet(path)
    if len(data) != manifest["rows"] or data.duplicated(GRAIN).any() or not data.frequency.eq("annual").all() or not data.index_measure.eq("same_year_previous_year_pct").all():
        raise ValueError("Invalid annual candidate grain or index basis.")
    sources = {s["source_file"]: s["sha256"] for s in manifest["sources"]}
    if not data.source_file.map(sources).eq(data.source_sha256).all():
        raise ValueError("Annual source hashes disagree.")
    numeric = data.index_value.dropna()
    if not np.isfinite(numeric).all() or numeric.lt(0).any() or not np.allclose(data.growth_yoy_pct, data.index_value - 100, equal_nan=True):
        raise ValueError("Invalid annual values or growth calculation.")
    data["period"] = pd.to_datetime(data.period, errors="raise")
    if data.period.isna().any() or not (data.period.dt.month.eq(12) & data.period.dt.day.eq(31)).all():
        raise ValueError("Invalid annual periods.")
    return data, manifest


def render_annual_preview(st, alt, root, region_id):
    st.subheader("Annual industrial production")
    try:
        data, manifest = load_annual_candidate(Path(root))
    except (OSError, ValueError, KeyError) as exc:
        st.error(f"Annual production withheld: {exc}")
        return
    st.info("Annual research review dataset · publication does not establish independent verification.")
    region = data.loc[data.region_id.eq(region_id)]
    section = st.selectbox("Annual production scope", list(s[0] for s in SHEETS.values()), format_func=lambda key: next(s[1] for s in SHEETS.values() if s[0] == key))
    shown = region.loc[region.okved_section.eq(section)].copy()
    if shown.empty:
        st.info("No eligible annual observations for this region.")
        return
    shown["Year"] = shown.period.dt.year.astype(str)
    st.caption("Published whole-year change versus the previous year · index minus 100 · not an average of monthly indices")
    st.warning("The latest workbook explicitly changes headline weights to 2023. Earlier observations come from a different vintage with an unstated weight base. Years are shown as separate bars, not spliced into a constant-method output level.")
    chart = alt.Chart(shown).mark_bar(color="#2455A4").encode(
        x=alt.X("Year:N", title=None, axis=alt.Axis(labelAngle=0)),
        y=alt.Y("growth_yoy_pct:Q", title="Annual YoY (%)", scale=alt.Scale(zero=True)),
        tooltip=["Year", alt.Tooltip("growth_yoy_pct:Q", title="Annual YoY (%)", format="+.1f"), "source_vintage", "index_base", "availability"])
    zero = alt.Chart(pd.DataFrame({"zero": [0]})).mark_rule(color="#182230").encode(y="zero:Q")
    if shown.index_value.notna().any():
        st.altair_chart((chart + zero).properties(height=310, description="Published annual production growth. Missing and suppressed observations remain unavailable."), width="stretch")
    else:
        st.info("No numeric annual indices are published for this region and scope. The original source markers remain in the table and download.")
    display = shown[["Year", "index_value", "growth_yoy_pct", "availability", "source_vintage", "index_base", "revision_state"]].copy()
    for column in ["index_value", "growth_yoy_pct"]:
        display[column] = display[column].map(lambda value: f"{value:+.1f}" if pd.notna(value) and column == "growth_yoy_pct" else f"{value:.1f}" if pd.notna(value) else "Unavailable")
    st.dataframe(display, hide_index=True, width="stretch")
    with st.expander("Annual sources and audit limits"):
        st.write("Sources: " + ", ".join(shown.source_file.unique()))
        for limitation in manifest["limitations"]:
            st.write(limitation)
        st.caption("Only the headline and broad OKVED2 sections B–E are available in these workbooks, not detailed production types. No national aggregate or causal attribution is calculated.")
    st.download_button("Download annual production evidence (CSV)", shown.to_csv(index=False).encode("utf-8-sig"), file_name=f"annual_production_{region_id}_{section}.csv", mime="text/csv")
