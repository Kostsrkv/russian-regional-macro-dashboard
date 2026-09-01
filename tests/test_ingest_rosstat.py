from pathlib import Path

from openpyxl import Workbook

from macro_rus.ingest_rosstat import SHEETS, ingest_rosstat_result


def _monthly_workbook(path: Path):
    workbook = Workbook()
    contents = workbook.active
    contents.title = "Содержание"
    contents["A1"] = "Обновлено: 26.08.2026 г."
    for sheet_name in [str(number) for number in range(1, 16)]:
        sheet = workbook.create_sheet(sheet_name)
        if sheet_name in SHEETS:
            sheet.cell(4, 2, "2026 год")
            sheet.cell(5, 2, "январь")
            sheet.cell(6, 1, "Калужская область")
            sheet.cell(6, 2, 101.5)
    workbook.save(path)


def test_rosstat_reads_monthly_headline_and_be_sections(tmp_path: Path):
    path = tmp_path / "ind_sub_2023_07-2026.xlsx"
    _monthly_workbook(path)
    result = ingest_rosstat_result(tmp_path)
    frame = result.industrial_production
    assert len(frame) == 5
    assert set(frame["okved_section"]) == {"TOTAL", "B", "C", "D", "E"}
    assert set(frame["index_base"]) == {"2023-base basket"}
    assert frame["source_id"].nunique() == 1

