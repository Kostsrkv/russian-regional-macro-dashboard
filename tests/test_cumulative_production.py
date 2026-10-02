"""Small source-shaped fixtures for cumulative basis, vintage and grain rules."""
from __future__ import annotations

import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from macro_rus.cumulative_production import (
    AUDIT_FILES, DATASET, GRAIN, SHEETS, TITLE_PHRASES, build_cumulative_candidate,
    load_cumulative_candidate, normalize, overlap_audit, read_cumulative,
    select_vintages, window_month,
)
from macro_rus.provenance import sha256_file


RECORDS = [
    dict(region_id="RU-KLU", treasury_name_ru="Калужская область", rosstat_name_ru="Калужская область", rosstat_match_method="typography_normalization"),
    dict(region_id="RU-SVE", treasury_name_ru="Свердловская область", rosstat_name_ru="Свердловская область", rosstat_match_method="typography_normalization"),
]


def workbook_fixture(path, date="01.04.2026", value=105.2, missing=False, basis=None, header=None, title=None, duplicate=False):
    workbook = Workbook()
    contents = workbook.active
    contents.title = "Содержание"
    contents["B1"] = f"Обновлено: {date} г."
    for name, (section, _, _) in SHEETS.items():
        sheet = workbook.create_sheet(name)
        sheet["A2"] = title or TITLE_PHRASES[section]
        sheet["A3"] = basis or "в % к соответствующему периоду предыдущего года"
        sheet["B4"] = "2026 год2"
        for col, window in enumerate(["январь", "январь-февраль", "январь-март3"], start=2):
            sheet.cell(5, col, header if col == 4 and header else window)
        for i, record in enumerate(RECORDS, start=6):
            sheet.cell(i, 1, record["rosstat_name_ru"])
            for col in range(2, 5):
                sheet.cell(i, col, value)
        if missing:
            sheet["B6"] = None
            sheet["C6"] = "…3"
            sheet["D6"] = 0
        if duplicate:
            sheet["A8"] = RECORDS[0]["rosstat_name_ru"]
            sheet["B8"] = value
        if section == "TOTAL":
            sheet["A10"] = "1 В качестве весов используется структура валовой добавленной стоимости 2023 базисного года."
    workbook.save(path)
    return path


def test_reads_published_cumulative_basis_and_unchanged_pilot_identifiers(tmp_path):
    data, source = read_cumulative(workbook_fixture(tmp_path / "source.xlsx"), RECORDS)
    assert len(data) == 30
    assert set(data.region_id) == {"RU-KLU", "RU-SVE"}
    assert not data.duplicated(GRAIN).any()
    assert data.frequency.eq("YTD").all()
    assert data.index_measure.eq("same_period_previous_year_pct").all()
    assert data.window_start.eq(pd.Timestamp("2026-01-01")).all()
    assert data.period.dt.is_month_end.all()
    q1 = data[data.period.eq(pd.Timestamp("2026-03-31"))]
    assert q1.raw_window_header.eq("январь-март3").all()
    assert q1.window_months.eq(3).all()
    assert q1.index_value.eq(105.2).all()
    assert source["source_vintage"] == "2026-04-01"
    assert data[data.okved_section.eq("TOTAL")].index_base.str.startswith("2023 GVA").all()
    assert data[data.okved_section.eq("B")].index_base.eq("Weight base not stated in this sheet").all()


@pytest.mark.parametrize("kwargs", [
    {"basis": "в % к соответствующему месяцу предыдущего года"},
    {"header": "март"}, {"header": "февраль-март"},
    {"header": "январь-апрель"}, {"title": "Неизвестный показатель"},
])
def test_rejects_wrong_title_basis_or_window(tmp_path, kwargs):
    with pytest.raises(ValueError):
        read_cumulative(workbook_fixture(tmp_path / "wrong.xlsx", **kwargs), RECORDS)


def test_keeps_blank_suppressed_and_real_zero_explicit(tmp_path):
    data, _ = read_cumulative(workbook_fixture(tmp_path / "missing.xlsx", missing=True), RECORDS)
    rows = data[data.region_id.eq("RU-KLU") & data.okved_section.eq("B")].sort_values("period")
    assert rows.availability.tolist() == ["unavailable", "unavailable_source_marker", "reported"]
    assert rows.raw_value.tolist() == ["", "…3", "0"]
    assert rows.index_value.isna().tolist() == [True, True, False]
    assert rows.index_value.iloc[2] == 0
    assert rows.growth_yoy_pct.iloc[2] == -100


def test_newest_adjacent_vintage_wins_even_missing_and_audit_retains_original(tmp_path):
    old, _ = read_cumulative(workbook_fixture(tmp_path / "old.xlsx", date="01.04.2026"), RECORDS)
    new, _ = read_cumulative(workbook_fixture(tmp_path / "new.xlsx", date="02.04.2026", missing=True, value=107), RECORDS)
    selected = select_vintages([new, old])
    assert len(selected) == 30
    assert selected.source_file.eq("new.xlsx").all()
    assert selected.loc[selected.region_id.eq("RU-KLU") & selected.window_months.eq(1), "index_value"].isna().all()
    audit = overlap_audit([old, new])
    assert len(audit) == 60
    assert audit.is_selected_vintage.sum() == 30
    assert len(audit[GRAIN].drop_duplicates()) == 30
    assert audit.loc[audit.source_file.eq("old.xlsx") & audit.window_months.eq(1) & audit.region_id.eq("RU-KLU"), "comparison_status"].eq("latest_missing").all()
    assert old.index_value.eq(105.2).all()


def test_duplicate_rows_dates_and_scope_qualifiers_fail_closed(tmp_path):
    with pytest.raises(ValueError, match="Duplicate regional"):
        read_cumulative(workbook_fixture(tmp_path / "duplicate.xlsx", duplicate=True), RECORDS)
    data, _ = read_cumulative(workbook_fixture(tmp_path / "source.xlsx"), RECORDS)
    with pytest.raises(ValueError, match="same-date"):
        select_vintages([data, data])
    assert normalize("Тюменская область") != normalize("Тюменская область без авт. округов")
    assert normalize("Чукотский авт. округ") == normalize("Чукотский авт.округ")


def test_window_footnotes_never_change_period():
    assert window_month("январь3") == 1
    assert window_month("январь–март3") == 3
    assert window_month("январь-декабрь") == 12
    with pytest.raises(ValueError):
        window_month("март")


def candidate_fixture(tmp_path):
    data, source = read_cumulative(workbook_fixture(tmp_path / "source.xlsx"), RECORDS)
    output = tmp_path / "candidate"
    output.mkdir()
    data.to_parquet(output / DATASET, index=False)
    # Artifact hashes cover the mandatory evidence set. Parsing is independently
    # exercised above; this fixture isolates loader data and manifest checks.
    for filename in AUDIT_FILES:
        data.to_csv(output / filename, index=False)
    manifest = dict(status="candidate_not_promoted", blocking_failures=0, rows=len(data),
                    sources=[source], eligible_region_ids=[r["region_id"] for r in RECORDS],
                    dataset_sha256=sha256_file(output / DATASET),
                    artifact_sha256={f: sha256_file(output / f) for f in AUDIT_FILES})
    (output / "manifest.json").write_text(json.dumps(manifest))
    return output, data, manifest


def test_loader_accepts_valid_candidate_and_rejects_changed_artifact(tmp_path):
    output, data, manifest = candidate_fixture(tmp_path)
    loaded, _ = load_cumulative_candidate(output)
    assert len(loaded) == len(data)
    (output / "coverage.csv").write_text("corrupt")
    with pytest.raises(ValueError, match="artifact integrity"):
        load_cumulative_candidate(output)


@pytest.mark.parametrize("mutation", ["frequency", "grain", "window", "vintage", "missing", "growth"])
def test_loader_rejects_semantic_corruption_even_with_updated_hash(tmp_path, mutation):
    output, data, manifest = candidate_fixture(tmp_path)
    if mutation == "frequency":
        data.loc[0, "frequency"] = "monthly"
    elif mutation == "grain":
        data.iloc[0] = data.iloc[1]
    elif mutation == "window":
        data.loc[0, "window_start"] = pd.Timestamp("2026-02-01")
    elif mutation == "vintage":
        data.loc[0, "source_vintage"] = "2026-03-31"
    elif mutation == "missing":
        data.loc[0, "index_value"] = np.nan
        data.loc[0, "growth_yoy_pct"] = np.nan
    else:
        data.loc[0, "growth_yoy_pct"] = 42
    data.to_parquet(output / DATASET, index=False)
    manifest["dataset_sha256"] = sha256_file(output / DATASET)
    (output / "manifest.json").write_text(json.dumps(manifest))
    with pytest.raises(ValueError):
        load_cumulative_candidate(output)


def test_builder_refuses_existing_directory_before_source_reads(tmp_path):
    with pytest.raises(FileExistsError):
        build_cumulative_candidate(tmp_path / "does-not-exist", tmp_path)
